"""Build site/products.json and site/images/ from data/products.xlsx + photos/.
Usage:  pip install -r requirements.txt  &&  python build.py
"""
import json, shutil, sys
from pathlib import Path
from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).parent
SITE = ROOT / "site"; IMG = SITE / "images"
PHOTO_W = 800
(IMG / "specs").mkdir(parents=True, exist_ok=True)

def find_photo(code):
    for ext in ("webp", "jpg", "jpeg", "png"):
        p = ROOT / "photos" / f"{code}.{ext}"
        if p.exists(): return p

def save_photo(src, dst):
    im = Image.open(src)
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA"); bg = Image.new("RGB", im.size, "white"); bg.paste(im, mask=im.split()[3]); im = bg
    im = im.convert("RGB")
    if im.width > PHOTO_W: im = im.resize((PHOTO_W, round(im.height * PHOTO_W / im.width)), Image.LANCZOS)
    im.save(dst, "WEBP", quality=82, method=6)

def s(v):
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v).strip()

ws = load_workbook(ROOT / "data" / "products.xlsx", data_only=True)["Products"]
out, missing = [], []
for row in ws.iter_rows(min_row=2, values_only=True):
    code, title, brand, cat, specs, active = (list(row) + [None] * 6)[:6]
    code = s(code)
    if not code or s(active).lower() == "no": continue
    photo = find_photo(code)
    if not photo: missing.append(code); continue
    save_photo(photo, IMG / f"{code}.webp")
    strip = ROOT / "strips" / f"{code}.webp"
    has_strip = strip.exists() and not s(specs)
    if has_strip: shutil.copy(strip, IMG / "specs" / f"{code}.webp")
    out.append(dict(code=code, title=s(title), brand=s(brand), category=s(cat) or "Other", specs=s(specs),
        img=f"images/{code}.webp", strip=f"images/specs/{code}.webp" if has_strip else "",
        search=" ".join([s(title), s(specs), s(brand), s(cat)]).lower()))
(SITE / "products.json").write_text(json.dumps(out, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
print(f"{len(out)} products written to site/products.json")
if missing: print("WARNING - no photo found in photos/ for codes:", ", ".join(missing), file=sys.stderr)
